"""Tests for TitlePagePlugin — rendering and Report integration."""
import pytest
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH

from docx_report_gen import Report
from docx_report_gen_title_page import TitlePageConfig, TitlePagePlugin


# ---------- construction ----------

def test_plugin_name():
    assert TitlePagePlugin.name == 'title-page'


def test_plugin_accepts_kwargs():
    p = TitlePagePlugin(institution='X', year=2026)
    assert p._initial_config is not None
    assert p._initial_config.institution == 'X'
    assert p._initial_config.year == 2026


def test_plugin_accepts_config_object():
    cfg = TitlePageConfig(institution='Y')
    p = TitlePagePlugin(config=cfg)
    assert p._initial_config is cfg


def test_plugin_rejects_config_and_kwargs_together():
    with pytest.raises(ValueError, match='not both'):
        TitlePagePlugin(config=TitlePageConfig(), institution='X')


def test_plugin_rejects_unknown_field():
    with pytest.raises(TypeError):
        TitlePagePlugin(nonexistent_field='X')


def test_plugin_without_arguments_does_not_auto_render():
    r = Report(plugins=[TitlePagePlugin()])
    non_empty = [p.text for p in r.doc.paragraphs if p.text]
    assert non_empty == []


# ---------- block methods ----------

def test_report_has_title_page_method():
    r = Report(plugins=[TitlePagePlugin()])
    assert hasattr(r, 'title_page')
    assert callable(r.title_page)


def test_report_has_load_title_page_method():
    r = Report(plugins=[TitlePagePlugin()])
    assert hasattr(r, 'load_title_page')
    assert callable(r.load_title_page)


def test_title_page_method_accepts_kwargs():
    r = Report(plugins=[TitlePagePlugin()])
    r.title_page(institution='СПбПУ', year=2026)
    texts = [p.text for p in r.doc.paragraphs if p.text]
    assert 'СПбПУ' in texts
    assert '2026' in texts


def test_title_page_method_accepts_config_object():
    r = Report(plugins=[TitlePagePlugin()])
    cfg = TitlePageConfig(institution='FromConfig')
    r.title_page(cfg)
    texts = [p.text for p in r.doc.paragraphs if p.text]
    assert 'FromConfig' in texts


def test_title_page_method_rejects_config_and_kwargs():
    r = Report(plugins=[TitlePagePlugin()])
    with pytest.raises(ValueError, match='not both'):
        r.title_page(TitlePageConfig(), institution='X')


def test_title_page_method_returns_report():
    r = Report(plugins=[TitlePagePlugin()])
    result = r.title_page(institution='X')
    assert result is r


def test_title_page_method_chains():
    r = Report(plugins=[TitlePagePlugin()])
    r.title_page(institution='X').h1('Body')
    texts = [p.text for p in r.doc.paragraphs if p.text]
    assert 'X' in texts
    assert 'Body' in texts


def test_title_page_method_can_be_called_multiple_times():
    r = Report(plugins=[TitlePagePlugin(page_break=False)])
    r.title_page(institution='First')
    r.title_page(institution='Second')
    texts = [p.text for p in r.doc.paragraphs if p.text]
    assert 'First' in texts
    assert 'Second' in texts


def test_title_page_method_page_break_override():
    import zipfile
    from lxml import etree

    r = Report(plugins=[TitlePagePlugin()])
    r.title_page(institution='X', page_break=False)

    path = r.doc
    # Save to temp and check XML
    import tempfile
    with tempfile.NamedTemporaryFile(suffix='.docx', delete=False) as f:
        path = f.name
    r.save(path)

    with zipfile.ZipFile(path) as zf:
        xml = etree.fromstring(zf.read('word/document.xml'))
    W = 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'
    breaks = [
        br for br in xml.iter(f'{{{W}}}br')
        if br.get(f'{{{W}}}type') == 'page'
    ]
    assert not breaks


# ---------- auto-render in setup ----------

def test_auto_render_from_kwargs():
    r = Report(plugins=[TitlePagePlugin(institution='Auto')])
    texts = [p.text for p in r.doc.paragraphs if p.text]
    assert 'Auto' in texts


def test_auto_render_from_config():
    cfg = TitlePageConfig(institution='AutoCfg', year=2026)
    r = Report(plugins=[TitlePagePlugin(config=cfg)])
    texts = [p.text for p in r.doc.paragraphs if p.text]
    assert 'AutoCfg' in texts
    assert '2026' in texts


def test_auto_render_adds_page_break_by_default():
    import zipfile
    from lxml import etree
    import tempfile

    r = Report(plugins=[TitlePagePlugin(institution='X')])
    with tempfile.NamedTemporaryFile(suffix='.docx', delete=False) as f:
        path = f.name
    r.save(path)

    with zipfile.ZipFile(path) as zf:
        xml = etree.fromstring(zf.read('word/document.xml'))
    W = 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'
    breaks = [
        br for br in xml.iter(f'{{{W}}}br')
        if br.get(f'{{{W}}}type') == 'page'
    ]
    assert breaks


# ---------- source mode ----------

def _create_title_docx(tmp_path):
    """Create a minimal title-page-like .docx for source-mode tests."""
    doc = Document()
    p1 = doc.add_paragraph()
    p1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p1.add_run('СПбПУ').bold = True

    p2 = doc.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p2.add_run('ЛАБОРАТОРНАЯ РАБОТА №2')

    path = tmp_path / 'title.docx'
    doc.save(str(path))
    return path


def test_auto_render_from_source(tmp_path):
    src = _create_title_docx(tmp_path)
    r = Report(plugins=[TitlePagePlugin(source=src)])
    texts = [p.text for p in r.doc.paragraphs if p.text]
    assert 'СПбПУ' in texts
    assert 'ЛАБОРАТОРНАЯ РАБОТА №2' in texts


def test_source_mode_rejects_config(tmp_path):
    src = _create_title_docx(tmp_path)
    with pytest.raises(ValueError, match='not both'):
        TitlePagePlugin(source=src, institution='X')


def test_source_mode_rejects_missing_file(tmp_path):
    with pytest.raises(FileNotFoundError):
        TitlePagePlugin(source=tmp_path / 'nope.docx')


def test_load_title_page_method(tmp_path):
    src = _create_title_docx(tmp_path)
    r = Report(plugins=[TitlePagePlugin()])
    r.load_title_page(src)
    texts = [p.text for p in r.doc.paragraphs if p.text]
    assert 'СПбПУ' in texts


def test_load_title_page_returns_report(tmp_path):
    src = _create_title_docx(tmp_path)
    r = Report(plugins=[TitlePagePlugin()])
    result = r.load_title_page(src)
    assert result is r


def test_load_title_page_chains(tmp_path):
    src = _create_title_docx(tmp_path)
    r = Report(plugins=[TitlePagePlugin()])
    r.load_title_page(src).h1('Body')
    texts = [p.text for p in r.doc.paragraphs if p.text]
    assert 'СПбПУ' in texts
    assert 'Body' in texts


def test_load_title_page_missing_file(tmp_path):
    r = Report(plugins=[TitlePagePlugin()])
    with pytest.raises(FileNotFoundError):
        r.load_title_page(tmp_path / 'nope.docx')


def test_load_title_page_page_break_override(tmp_path):
    import zipfile
    from lxml import etree

    src = _create_title_docx(tmp_path)
    r = Report(plugins=[TitlePagePlugin()])
    r.load_title_page(src, page_break=False)

    path = tmp_path / 'out.docx'
    r.save(str(path))

    with zipfile.ZipFile(path) as zf:
        xml = etree.fromstring(zf.read('word/document.xml'))
    W = 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'
    breaks = [
        br for br in xml.iter(f'{{{W}}}br')
        if br.get(f'{{{W}}}type') == 'page'
    ]
    assert not breaks


def test_source_mode_preserves_alignment(tmp_path):
    src = _create_title_docx(tmp_path)
    r = Report(plugins=[TitlePagePlugin(source=src)])
    for para in r.doc.paragraphs:
        if para.text == 'СПбПУ':
            assert para.alignment == WD_ALIGN_PARAGRAPH.CENTER
            break
    else:
        pytest.fail('СПбПУ not found')


def test_source_mode_user_content_after_title(tmp_path):
    src = _create_title_docx(tmp_path)
    r = Report(plugins=[TitlePagePlugin(source=src)])
    r.h1('Body heading')
    texts = [p.text for p in r.doc.paragraphs if p.text]
    assert texts.index('СПбПУ') < texts.index('Body heading')


# ---------- global registration ----------

def test_global_registration_gives_methods_to_new_reports():
    from docx_report_gen.plugins import plugins as global_plugins

    plugin = TitlePagePlugin()
    try:
        global_plugins.register(plugin)
        r = Report()
        assert hasattr(r, 'title_page')
        r.title_page(institution='Global')
        texts = [p.text for p in r.doc.paragraphs if p.text]
        assert 'Global' in texts
    finally:
        global_plugins.unregister(plugin)


# ---------- config mode full render (kept from before) ----------

def test_full_title_page(tmp_path):
    r = Report(plugins=[TitlePagePlugin(
        institution='СПбПУ',
        institute='ИКНТ',
        school='ВШПИ',
        work_type='ЛАБОРАТОРНАЯ РАБОТА',
        work_number='№2',
        work_title='«Разложение в ряд Фурье»',
        discipline='«Цифровая обработка сигналов»',
        student='студент гр. 101 Ерохин В.С.',
        supervisor='Тутыгин В.С.',
        city='Санкт-Петербург',
        year=2026,
    )])
    r.h1('Introduction')
    path = tmp_path / 'out.docx'
    r.save(str(path))

    doc = Document(str(path))
    texts = [p.text for p in doc.paragraphs]
    assert 'СПбПУ' in texts
    assert 'ИКНТ' in texts
    assert 'ВШПИ' in texts
    assert 'ЛАБОРАТОРНАЯ РАБОТА №2' in texts
    assert '«Разложение в ряд Фурье»' in texts
    assert 'по дисциплине' in texts
    assert '«Цифровая обработка сигналов»' in texts
    assert 'Выполнил' in texts
    assert 'студент гр. 101 Ерохин В.С.' in texts
    assert 'Руководитель Тутыгин В.С.' in texts
    assert 'Санкт-Петербург' in texts
    assert '2026' in texts
    assert 'Introduction' in texts


def test_missing_fields_are_skipped(tmp_path):
    r = Report(plugins=[TitlePagePlugin(institution='X', city='Y')])
    r.h1('Body')
    path = tmp_path / 'out.docx'
    r.save(str(path))

    doc = Document(str(path))
    texts = [t for t in (p.text for p in doc.paragraphs) if t]
    assert 'X' in texts
    assert 'Y' in texts
    assert 'Body' in texts


def test_empty_config_produces_only_page_break():
    r = Report(plugins=[TitlePagePlugin()])
    r.title_page()
    non_empty = [t for t in (p.text for p in r.doc.paragraphs) if t]
    assert non_empty == []